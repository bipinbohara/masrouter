import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from MAR.Experiment.trace import collect_traces, run_graph
from MAR.Experiment.results import ResultWriter, parse_mmlu_answer
from MAR.Experiment.mmlu import evaluate


class FakeLLM:
    model_name = 'local/model'
    def gen(self, messages, **kwargs):
        return r'The answer is \boxed{A}.'
    async def agen(self, messages, **kwargs):
        return self.gen(messages, **kwargs)


class FakeGraph:
    reasoning_name = 'Debate'
    def __init__(self):
        self.nodes = {'one': SimpleNamespace(id='one', role=SimpleNamespace(role='Scientist'),
                                            llm=FakeLLM(), outputs=[])}
        self.decision_node = SimpleNamespace(id='final', llm=FakeLLM(), outputs=[])
    def run(self, inputs, num_rounds):
        for _ in range(num_rounds):
            node = self.nodes['one']
            node.outputs.append(node.llm.gen(inputs['query']))
        self.decision_node.outputs.append(self.decision_node.llm.gen(inputs['query']))
        return self.decision_node.outputs, 0


class TraceTests(unittest.TestCase):
    def test_opt_in_observation_counts_rounds_and_final_without_altering_output(self):
        graph = FakeGraph()
        expected = FakeGraph().run({'query': 'question'}, 3)
        self.assertEqual(run_graph(graph, {'query': 'question'}, 3), expected)
        graph = FakeGraph()
        with collect_traces() as traces:
            result = run_graph(graph, {'query': 'question'}, 3)
        self.assertEqual(result, expected)
        self.assertEqual(len(traces[0]['calls']), 4)
        self.assertEqual([c['agent_id'] for c in traces[0]['calls']], ['one'] * 3 + ['final'])
        self.assertEqual(len(traces[0]['agents']), 2)
        self.assertNotIn('gen', vars(graph.decision_node.llm))
        self.assertTrue(all(c['usage'] is None for c in traces[0]['calls']))
        with collect_traces() as other:
            pass
        self.assertEqual(other, [])

    def test_call_error_and_method_restoration(self):
        graph = FakeGraph()
        def fail(messages):
            raise RuntimeError('server unavailable')
        graph.decision_node.llm.gen = fail
        with collect_traces() as traces:
            with self.assertRaises(RuntimeError):
                run_graph(graph, {'query': 'question'}, 1)
        self.assertEqual(traces[0]['status'], 'error')
        self.assertEqual(traces[0]['calls'][-1]['status'], 'error')
        self.assertIs(graph.decision_node.llm.gen, fail)
        self.assertNotIn('gen', vars(graph.nodes['one'].llm))


class FakeDataset:
    split = 'test'
    def __init__(self):
        self.rows = [dict(question=f'Question {i}', A='a', B='b', C='c', D='d',
                          correct_answer='A', subject='biology', source_file='biology_test.csv', source_row=i)
                     for i in range(3)]
    def __len__(self):
        return len(self.rows)
    def __getitem__(self, index):
        return self.rows[index]
    def record_to_input(self, row):
        return {'task': row['question']}
    def record_to_target_answer(self, row):
        return row['correct_answer']


class FakeRouter:
    def __init__(self, fail_index=None):
        self.seen = []
        self.fail_index = fail_index
    def forward(self, queries, tasks, llms, reasonings, labels, prompt_file):
        self.seen.extend(queries)
        if len(self.seen) - 1 == self.fail_index:
            raise RuntimeError('routing failed')
        result = run_graph(FakeGraph(), {'query': queries[0]}, 2)
        return [result[0][0]], [0.0], None, None, None, [[1.7]]


class ResultTests(unittest.TestCase):
    def test_excel_path_cannot_overwrite_lossless_json(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, '.json'):
                ResultWriter(Path(directory) / 'run.xlsx', {})

    def test_keyboard_interrupt_saves_in_progress_task(self):
        class InterruptedRouter:
            def forward(self, *args, **kwargs):
                raise KeyboardInterrupt
        with tempfile.TemporaryDirectory() as directory:
            writer = ResultWriter(Path(directory) / 'run.json', {})
            with self.assertRaises(KeyboardInterrupt):
                evaluate(InterruptedRouter(), FakeDataset(), writer, [], [], [], 'prompt')
            writer.finish('interrupted')
            record = json.loads(writer.path.read_text())['tasks'][0]
            self.assertEqual(record['status'], 'error')
            self.assertIn('interrupted', record['error'])
    def test_evaluation_rejects_dev_split(self):
        dataset = FakeDataset()
        dataset.split = 'dev'
        with self.assertRaisesRegex(ValueError, 'test split'):
            evaluate(FakeRouter(), dataset, None, [], [], [], 'prompt')
    def test_parser_uses_final_answer_and_rejects_unstructured_prose(self):
        self.assertEqual(parse_mmlu_answer(r'Consider B and C. Final answer: \boxed{D}'), 'D')
        self.assertEqual(parse_mmlu_answer('Answer is (b).'), 'B')
        self.assertEqual(parse_mmlu_answer('C'), 'C')
        self.assertIsNone(parse_mmlu_answer('A could be right, or perhaps B'))

    def test_entire_test_split_saved_with_failure_and_no_skipped_first_or_last_row(self):
        with tempfile.TemporaryDirectory() as directory:
            writer = ResultWriter(Path(directory) / 'experiment.json', {'split': 'test'})
            router = FakeRouter(fail_index=1)
            evaluate(router, FakeDataset(), writer, [], [], [], 'prompt')
            # Records already exist before final consolidation.
            records = [json.loads(line) for line in writer.records_path.read_text().splitlines()]
            self.assertEqual(len(records), 3)
            self.assertEqual([r['source_row'] for r in records], [0, 1, 2])
            self.assertEqual(records[1]['status'], 'error')
            self.assertEqual(records[0]['num_agents'], 1)
            self.assertEqual(records[0]['num_nodes_including_final'], 2)
            self.assertEqual(records[0]['total_llm_calls'], 3)
            self.assertEqual(records[0]['predicted_agent_count'], 1.7)
            writer.finish(excel=True)
            report = json.loads(writer.path.read_text())
            self.assertEqual(report['summary']['accuracy'], 2/3)
            self.assertEqual(report['summary']['total_llm_calls'], 6)
            self.assertEqual(report['summary']['failed'], 1)
            self.assertEqual(report['summary']['subjects']['biology']['total'], 3)
            self.assertTrue(writer.path.with_suffix('.xlsx').exists())
            import openpyxl
            workbook = openpyxl.load_workbook(writer.path.with_suffix('.xlsx'))
            self.assertEqual(workbook['tasks'].max_row, 4)
            self.assertEqual(workbook['agents'].max_row, 5)
            self.assertEqual(workbook['calls_1'].max_row, 7)
            with self.assertRaises(FileExistsError):
                ResultWriter(writer.path, {})

    def test_limit_and_interrupted_export(self):
        with tempfile.TemporaryDirectory() as directory:
            writer = ResultWriter(Path(directory) / 'run.json', {})
            router = FakeRouter()
            evaluate(router, FakeDataset(), writer, [], [], [], 'prompt', max_tasks=1)
            writer.finish('interrupted')
            result = json.loads(writer.path.read_text())
            self.assertEqual(len(router.seen), 1)
            self.assertEqual(len(result['tasks']), 1)
            self.assertEqual(result['status'], 'interrupted')

    def test_excel_limits_and_formula_escaping_keep_json_lossless(self):
        with tempfile.TemporaryDirectory() as directory:
            writer = ResultWriter(Path(directory) / 'run.json', {})
            query = '=malicious' + 'x' * 40000
            dataset = FakeDataset()
            dataset.rows[0]['question'] = query
            evaluate(FakeRouter(), dataset, writer, [], [], [], 'prompt', max_tasks=1)
            writer.finish(excel=True)
            self.assertEqual(json.loads(writer.path.read_text())['tasks'][0]['task'], query)
            import openpyxl
            sheet = openpyxl.load_workbook(writer.path.with_suffix('.xlsx'))['tasks']
            column = [cell.value for cell in sheet[1]].index('task') + 1
            cell = sheet.cell(2, column)
            self.assertEqual(cell.data_type, 's')
            self.assertLessEqual(len(cell.value), 32767)
            self.assertIn('TRUNCATED', cell.value)


class DatasetTests(unittest.TestCase):
    def test_mmlu_source_ids_survive_shuffle(self):
        from Datasets.mmlu_dataset import MMLUDataset
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / 'test'
            folder.mkdir()
            (folder / 'biology_test.csv').write_text('q1,a,b,c,d,A\nq2,a,b,c,d,B\n')
            rows = MMLUDataset._load_data(folder)
            self.assertEqual(len(rows), 2)
            for _, row in rows.iterrows():
                self.assertEqual(row['subject'], 'biology')
                self.assertEqual(row['source_file'], 'biology_test.csv')
                self.assertEqual(row['question'], f"q{int(row['source_row'])+1}")


class RunnerArgumentTests(unittest.TestCase):
    def parse(self, arguments):
        # Exercise the real CLI parser without importing GPU/model dependencies.
        import argparse
        import ast
        import sys
        from unittest.mock import patch
        source = Path('Experiments/run_mmlu.py').read_text()
        function = next(node for node in ast.parse(source).body
                        if isinstance(node, ast.FunctionDef) and node.name == 'parse_args')
        namespace = {'argparse': argparse}
        exec(compile(ast.Module(body=[function], type_ignores=[]), 'run_mmlu.py', 'exec'), namespace)
        with patch.object(sys, 'argv', ['run_mmlu.py', *arguments]):
            return namespace['parse_args']()

    def test_checkpoint_defaults_to_full_test_evaluation(self):
        args = self.parse(['--checkpoint', 'router.pth'])
        self.assertEqual(args.mode, 'test')
        self.assertIsNone(args.max_tasks)

    def test_untrained_evaluation_requires_explicit_acknowledgement(self):
        from contextlib import redirect_stderr
        import io
        with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            self.parse([])
        self.assertTrue(self.parse(['--allow_untrained']).allow_untrained)

    def test_train_test_remains_available_but_invalid_limits_fail(self):
        self.assertEqual(self.parse(['--mode', 'train-test']).mode, 'train-test')
        from contextlib import redirect_stderr
        import io
        with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            self.parse(['--allow_untrained', '--max_tasks', '0'])


if __name__ == '__main__':
    unittest.main()
